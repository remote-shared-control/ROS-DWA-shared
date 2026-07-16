#!/usr/bin/env python3
# -*- coding:utf-8 -*-

import os
import math

import rosbag
from openpyxl import load_workbook


# ===========================
# 路径
# ===========================
excel_path = "/home/frienkie/result/data.xlsx"
bag_dir = "/home/frienkie/rosbag"


# ===========================
# 计算rosbag中的运动距离、时间、平均速度
# ===========================
def calc_bag_info(bag_path):

    bag = rosbag.Bag(bag_path)

    last_x = None
    last_y = None

    start_time = None
    end_time = None

    distance = 0.0

    for topic, msg, t in bag.read_messages(topics=["/odom"]):

        x = msg.pose.pose.position.x
        y = msg.pose.pose.position.y

        current_time = t.to_sec()

        if start_time is None:
            start_time = current_time

        end_time = current_time

        if last_x is not None:
            dx = x - last_x
            dy = y - last_y
            distance += math.hypot(dx, dy)

        last_x = x
        last_y = y

    bag.close()

    if start_time is None or end_time is None:
        return None

    duration = end_time - start_time

    if duration > 0:
        avg_speed = distance / duration
    else:
        avg_speed = 0.0

    return distance, duration, avg_speed


# ===========================
# 主程序
# ===========================
def main():

    wb = load_workbook(excel_path)
    ws = wb.active

    # 从第二行开始（第一行为表头）
    for row in range(2, ws.max_row + 1):

        value = ws.cell(row=row, column=1).value

        if value is None:
            continue

        try:
            n = int(value)
        except:
            print(f"第{row}行A列不是数字，跳过")
            continue

        bag_path = os.path.join(bag_dir, f"rosbag{n}.bag")

        if not os.path.exists(bag_path):
            print(f"不存在: {bag_path}")
            continue

        print(f"处理 {bag_path}")

        result = calc_bag_info(bag_path)

        if result is None:
            print("没有odom数据")
            continue

        distance, duration, avg_speed = result

        # C D E列
        ws.cell(row=row, column=3).value = round(distance, 2)
        ws.cell(row=row, column=4).value = round(duration, 2)
        ws.cell(row=row, column=5).value = round(avg_speed, 2)

        print(
            f"距离={distance:.3f} m  时间={duration:.3f} s  平均速度={avg_speed:.3f} m/s"
        )

    wb.save(excel_path)
    print("完成！")


if __name__ == "__main__":
    main()