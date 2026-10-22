"""Check direct rosbag launch, independent coordinates, and trial statistics."""

import ast
import json
import math
import os
from pathlib import Path
import signal
import tempfile
import unittest
from types import SimpleNamespace
from unittest.mock import Mock, patch

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    Workbook = load_workbook = None


SIM_DIR = Path(__file__).resolve().parents[1]


def load_helpers():
    path = SIM_DIR / "distancetime.py"
    tree = ast.parse(path.read_text())
    names = {
        "get_time", "odom_callback", "record_robot_position",
        "get_next_rosbag_number", "start_rosbag", "stop_rosbag",
        "save_robot_coordinates", "save_result",
    }
    functions = [node for node in tree.body
                 if isinstance(node, ast.FunctionDef) and node.name in names]
    namespace = {
        "math": math, "os": os, "json": json, "signal": signal,
        "subprocess": Mock(), "rospy": Mock(), "rosbag_process": None,
        "Workbook": Workbook, "load_workbook": load_workbook,
        "print": lambda *a, **kw: None,
    }
    exec(compile(ast.Module(body=functions, type_ignores=[]), str(path), "exec"), namespace)
    return namespace


def config():
    return SimpleNamespace(x=0.0, y=0.0, th=0.0, v=0.0, w=0.0, xy=[],
                           first_time=True, prev_x=0.0, prev_y=0.0,
                           prev_x0=0.0, prev_y0=0.0, distance=0.0, distance0=0.0)


def odom(x, y):
    return SimpleNamespace(
        pose=SimpleNamespace(pose=SimpleNamespace(
            position=SimpleNamespace(x=x, y=y),
            orientation=SimpleNamespace(x=0, y=0, z=0, w=1))),
        twist=SimpleNamespace(twist=SimpleNamespace(
            linear=SimpleNamespace(x=0.0), angular=SimpleNamespace(z=0.0))),
    )


class DirectRecordingTests(unittest.TestCase):
    def test_launch_returns_after_popen_without_time_or_topic_checks(self):
        namespace = load_helpers()
        namespace["get_next_rosbag_number"] = lambda: 7
        with tempfile.TemporaryDirectory() as directory:
            # Existing output files must not cause numbering changes or refusal.
            for name in ("rosbag7.bag", "rosbag7.bag.active"):
                (Path(directory) / name).write_text("previous run")
            with patch.object(os.path, "expanduser", return_value=directory):
                result = namespace["start_rosbag"]()
            self.assertEqual(result, 7)
            namespace["subprocess"].Popen.assert_called_once_with([
                "rosbag", "record", "-O", str(Path(directory) / "rosbag7"),
                "/cmd_vel", "/cmd_vel_human", "/odom", "/dynamic_obstacle_states",
            ])
            process = namespace["subprocess"].Popen.return_value
            process.wait.assert_not_called()
            process.poll.assert_not_called()
            namespace["rospy"].get_param.assert_not_called()
            namespace["rospy"].Publisher.assert_not_called()
            namespace["rospy"].Subscriber.assert_not_called()
            namespace["rospy"].Time.now.assert_not_called()

    def test_stop_closes_standard_recorder_and_is_idempotent(self):
        namespace = load_helpers()
        process = Mock()
        namespace["rosbag_process"] = process
        namespace["stop_rosbag"]()
        namespace["stop_rosbag"]()
        process.send_signal.assert_called_once_with(signal.SIGINT)
        process.wait.assert_called_once_with()
        self.assertIsNone(namespace["rosbag_process"])


class CoordinateTests(unittest.TestCase):
    def test_initial_position_is_not_counted_as_distance_from_zero(self):
        namespace = load_helpers()
        robot = config()
        robot.x, robot.y = -3.6, -3.0
        namespace["record_robot_position"](robot)
        self.assertEqual(robot.distance, 0.0)
        self.assertEqual(robot.xy, [(-3.6, -3.0)])
        self.assertEqual((robot.prev_x, robot.prev_y), (-3.6, -3.0))

    def test_distance_and_coordinate_order_preserve_return_visits(self):
        namespace = load_helpers()
        robot = config()
        for x, y in ((0.0, 1.5), (0.3, 1.9), (0.3, 1.9), (0.0, 1.5)):
            robot.x, robot.y = x, y
            namespace["record_robot_position"](robot)
        self.assertAlmostEqual(robot.distance, 1.0)
        self.assertEqual(robot.xy, [(0.0, 1.5), (0.3, 1.9), (0.0, 1.5)])

    def test_submillimetre_noise_uses_original_distance_threshold(self):
        namespace = load_helpers()
        robot = config()
        namespace["record_robot_position"](robot)
        robot.x = 0.00001
        namespace["record_robot_position"](robot)
        self.assertEqual(robot.distance, 0.0)

    def test_coordinate_file_is_ordered_json_and_overwrites_same_id(self):
        namespace = load_helpers()
        points = [(0, 0), (1, 2), (0, 0)]
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "test7.txt"
            target.write_text("previous coordinates")
            output = namespace["save_robot_coordinates"](7, points, directory)
            self.assertEqual(output, str(target))
            self.assertEqual(json.loads(target.read_text()), [[0, 0], [1, 2], [0, 0]])

    def test_every_dwa_odom_callback_collects_coordinates_and_distance(self):
        scripts = sorted(SIM_DIR.glob("dwa*.py"))
        self.assertEqual(len(scripts), 9)
        for path in scripts:
            with self.subTest(script=path.name):
                namespace = load_helpers()
                namespace["euler_from_quaternion"] = lambda _: (0.0, 0.0, 0.0)
                tree = ast.parse(path.read_text())
                cls = next(node for node in tree.body
                           if isinstance(node, ast.ClassDef) and node.name == "Config")
                callback = next(node for node in cls.body if isinstance(node, ast.FunctionDef)
                                and node.name == "assignOdomCoords")
                exec(compile(ast.Module(body=[callback], type_ignores=[]), str(path), "exec"), namespace)
                robot = config()
                namespace["assignOdomCoords"](robot, odom(0.0, 1.5))
                namespace["assignOdomCoords"](robot, odom(0.3, 1.9))
                self.assertEqual(robot.xy, [(0.0, 1.5), (0.3, 1.9)])
                self.assertAlmostEqual(robot.distance, 0.5)


@unittest.skipUnless(Workbook is not None, "openpyxl is required for Excel persistence tests")
class ResultTests(unittest.TestCase):
    def test_new_excel_saves_duration_distance_and_separate_coordinates(self):
        namespace = load_helpers()
        namespace["rospy"].get_time.return_value = 100.678
        with tempfile.TemporaryDirectory() as directory:
            excel = Path(directory) / "results/data.xlsx"
            original_save = namespace["save_robot_coordinates"]
            namespace["save_robot_coordinates"] = lambda number, points: original_save(
                number, points, str(Path(directory) / "cood"))
            namespace["save_result"](
                7, -90.0, 2, file_path=str(excel), start_time=90.0,
                distance=1.2345, coordinates=[(0, 0), (0.3, 0.4)],
            )
            wb = load_workbook(excel, data_only=True)
            try:
                ws = wb.active
                self.assertEqual([ws.cell(1, col).value for col in (3, 4)],
                                 ["elapsed_time_s", "distance_m"])
                self.assertEqual([ws.cell(2, col).value for col in range(1, 7)],
                                 [7, 0, 10.68, 1.23, None, 2])
            finally:
                wb.close()
            points = json.loads((Path(directory) / "cood/test7.txt").read_text())
            self.assertEqual(points, [[0, 0], [0.3, 0.4]])

    def test_existing_excel_data_and_original_columns_are_preserved(self):
        namespace = load_helpers()
        namespace["rospy"].get_time.return_value = 20.0
        with tempfile.TemporaryDirectory() as directory:
            excel = Path(directory) / "data.xlsx"
            wb = Workbook()
            wb.active.append(["A", "B", "C", "D", "E", "F"])
            old_row = [3, 0.5, None, None, "keep", 1]
            wb.active.append(old_row)
            wb.save(excel)
            wb.close()
            namespace["save_result"](4, 3.0, 3, file_path=str(excel),
                                     start_time=10.0, distance=2.5)
            # Existing callers that only supply A/B/F remain valid.
            namespace["save_result"](5, 0.0, 0, file_path=str(excel))
            wb = load_workbook(excel, data_only=True)
            try:
                ws = wb.active
                self.assertEqual([ws.cell(2, col).value for col in range(1, 7)], old_row)
                self.assertEqual([ws.cell(3, col).value for col in range(1, 7)],
                                 [4, 3.0, 10.0, 2.5, None, 3])
                self.assertEqual([ws.cell(4, col).value for col in range(1, 7)],
                                 [5, 0.0, None, None, None, 0])
            finally:
                wb.close()
            self.assertEqual(namespace["get_next_rosbag_number"](str(excel)), 6)


class ControllerIntegrationTests(unittest.TestCase):
    def test_all_dwa_scripts_launch_directly_and_save_statistics_before_stop(self):
        for path in sorted(SIM_DIR.glob("dwa*.py")):
            with self.subTest(script=path.name):
                tree = ast.parse(path.read_text())
                main = next(node for node in tree.body
                            if isinstance(node, ast.FunctionDef) and node.name == "main")
                calls = [node for node in ast.walk(main) if isinstance(node, ast.Call)]
                starts = [call for call in calls if isinstance(call.func, ast.Name)
                          and call.func.id == "start_rosbag"]
                self.assertEqual(len(starts), 1)
                self.assertFalse(starts[0].args or starts[0].keywords)
                self.assertTrue(any(isinstance(call.func, ast.Attribute)
                                    and isinstance(call.func.value, ast.Name)
                                    and call.func.value.id == "pub" and call.func.attr == "publish"
                                    for call in calls))
                saves = [call for call in calls if isinstance(call.func, ast.Name)
                         and call.func.id == "save_result"]
                self.assertEqual(len(saves), 1 if "newturtle" in path.name else 2)
                for call in saves:
                    self.assertEqual({kw.arg for kw in call.keywords},
                                     {"start_time", "distance", "coordinates"})
                for node in ast.walk(main):
                    if not isinstance(node, ast.If):
                        continue
                    operations = [stmt.value.func.id for stmt in node.body
                                  if isinstance(stmt, ast.Expr) and isinstance(stmt.value, ast.Call)
                                  and isinstance(stmt.value.func, ast.Name)]
                    if "save_result" in operations:
                        self.assertEqual(operations, ["save_result", "stop_rosbag"])


if __name__ == "__main__":
    unittest.main()
